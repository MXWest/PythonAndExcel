from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.drawing.image import Image
from openpyxl import load_workbook

wb = load_workbook("data/Pie.xlsx")
ws = wb.active

table = Table(displayName="Table_1", ref="A1:B5")
style = TableStyleInfo(
    name="TableStyleMedium9",
    showFirstColumn=False,
    showLastColumn=False,
    showRowStripes=True,
    showColumnStripes=True,
)
table.tableStyleInfo = style
ws.add_table(table)
wb.save("data/mike_table.xlsx")

image = Image("data/madecraft.jpg")
image.height = image.height * 0.25
image.width = image.width * 0.25
ws.add_image(image, "C1")
wb.save("data/mike_table_image_crop.xlsx")
